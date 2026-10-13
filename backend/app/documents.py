"""Turns an upload or pasted text into a ResumeSource, checking type and size first.

The file type is read from the file's first bytes, never from its name or the
browser's Content-Type, both of which the uploader controls.
"""

import io
import zipfile
from dataclasses import dataclass

import docx
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError

from app.schemas import SourceType

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 30_000
# A .docx is a zip, and python-docx inflates every part it can reach (styles, headers,
# media), so a small file can declare gigabytes (a "zip bomb"). Declared sizes are
# trustworthy because zipfile stops reading at them; refuse before anything inflates.
MAX_DOCX_TOTAL_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 1000
DOCX_BODY = "word/document.xml"

ZIP_MAGIC = b"PK\x03\x04"
IMAGE_SIGNATURES: tuple[tuple[bytes, str], ...] = (
    (b"\x89PNG\r\n\x1a\n", "image/png"),
    (b"\xff\xd8\xff", "image/jpeg"),
)


@dataclass(frozen=True)
class FileBlob:
    data: bytes
    mime_type: str


@dataclass(frozen=True)
class ResumeSource:
    source_type: SourceType
    text: str | None
    blob: FileBlob | None


class InputProblem(Exception):
    status_code = 422
    user_message = "That input could not be used."


class EmptyInput(InputProblem):
    user_message = "The resume is empty. Upload a file or paste the resume text."


class TextTooLong(InputProblem):
    status_code = 413
    user_message = f"Pasted text is limited to {MAX_TEXT_CHARS:,} characters."


class FileTooLarge(InputProblem):
    status_code = 413
    user_message = f"Files are limited to {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."


class UnsupportedFile(InputProblem):
    status_code = 415
    user_message = "Upload a PDF, Word (.docx), PNG, JPG or WebP file."


class UnreadableDocument(InputProblem):
    user_message = "This Word file could not be read. Save it again, or upload it as a PDF."


class AmbiguousInput(InputProblem):
    user_message = "Send either a file or pasted text, not both."


def from_text(raw: str) -> ResumeSource:
    text = raw.strip()
    if not text:
        raise EmptyInput()
    if len(text) > MAX_TEXT_CHARS:
        raise TextTooLong()
    return ResumeSource(source_type="text", text=text, blob=None)


def from_upload(data: bytes) -> ResumeSource:
    if not data:
        raise EmptyInput()
    if len(data) > MAX_UPLOAD_BYTES:
        raise FileTooLarge()
    if data.startswith(b"%PDF-"):
        return ResumeSource(source_type="pdf", text=None, blob=FileBlob(data, "application/pdf"))
    image_mime = _image_mime(data)
    if image_mime:
        return ResumeSource(source_type="image", text=None, blob=FileBlob(data, image_mime))
    if data.startswith(ZIP_MAGIC):
        return ResumeSource(source_type="docx", text=_docx_text(data), blob=None)
    raise UnsupportedFile()


def _image_mime(data: bytes) -> str | None:
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    for signature, mime in IMAGE_SIGNATURES:
        if data.startswith(signature):
            return mime
    return None


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise UnreadableDocument() from exc
    if DOCX_BODY not in {entry.filename for entry in entries}:
        raise UnsupportedFile()  # a zip, but not a Word document
    if len(entries) > MAX_DOCX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_DOCX_TOTAL_BYTES:
        raise UnreadableDocument()

    try:
        text = _paragraphs_and_tables(data)
    except (PackageNotFoundError, XMLSyntaxError) as exc:
        raise UnreadableDocument() from exc
    except Exception as exc:  # untrusted-parser boundary: any failure means "unreadable", never a 500
        raise UnreadableDocument() from exc
    if not text:
        raise EmptyInput()
    if len(text) > MAX_TEXT_CHARS:
        raise TextTooLong()
    return text


def _paragraphs_and_tables(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())
