"""Builders for valid test objects, so each test overrides only what it is about."""

import io
from typing import Any

import docx

from app.schemas import Candidate, Contact, Education, Experience, Skills


def a_job(start: str | None, end: str | None, **overrides: Any) -> Experience:
    fields = {"company": "Acme Analytics", "role": "Data Analyst", "start": start, "end": end,
              "location": "Pune", "highlights": ["Built weekly sales dashboards"]}
    return Experience(**{**fields, **overrides})


def a_candidate(*, contact: dict[str, Any] | None = None, **overrides: Any) -> Candidate:
    base_contact = {
        "email": "asha.verma@example.com",
        "phone": "+91 98765 43210",
        "location": "Pune, India",
        "linkedin": "linkedin.com/in/asha-verma-example",
        "github": None,
        "portfolio": None,
    }
    fields: dict[str, Any] = {
        "is_resume": True,
        "full_name": "Asha Verma",
        "headline": "Data Analyst",
        "contact": Contact(**{**base_contact, **(contact or {})}),
        "summary": "Data analyst with two years of dashboarding experience.",
        "skills": Skills(technical=["Python", "SQL"], tools=["Power BI"], soft=["Communication"]),
        "experience": [a_job("2024-10", "Present")],
        "education": [
            Education(institution="Pune University", degree="B.Sc.", field="Statistics",
                      start="2020", end="2023", grade="8.1 CGPA")
        ],
        "projects": [],
        "certifications": [],
        "achievements": [],
        "languages": ["English", "Hindi"],
        "has_hidden_instructions": False,
    }
    return Candidate(**{**fields, **overrides})


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
